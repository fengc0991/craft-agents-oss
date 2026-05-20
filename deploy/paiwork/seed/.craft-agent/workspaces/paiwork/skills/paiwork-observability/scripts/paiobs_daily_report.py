#!/usr/bin/env python3.11
"""Orchestrate the PaiWork daily report workflow.

This script is intentionally scheduler-agnostic. The agent/platform scheduler
should invoke it at 01:00; by default it analyzes the previous local calendar
day from 00:00:00 to 00:00:00.
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import subprocess
import sys
import threading
import time
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any
from urllib.parse import urlencode
from zoneinfo import ZoneInfo

import paiobs
from paiobs_ai_analysis import (
    DEFAULT_AI_WORKERS,
    QUERY_MAJOR_CATEGORY_LABELS,
    QUERY_MINOR_CATEGORY_LABELS,
    chinese_label,
    issue_level1_for_level2,
    issue_owner_for_level2,
    normalize_issue_level2,
)

try:
    from openpyxl import Workbook
except ImportError:  # pragma: no cover - deployment image normally includes openpyxl.
    Workbook = None


SCRIPT_DIR = Path(__file__).resolve().parent
SKILL_DIR = SCRIPT_DIR.parent
DEFAULT_TIMEZONE = os.environ.get("PAI_OBS_DAILY_TIMEZONE", "Asia/Shanghai")
DEFAULT_OUTPUT_DIR = os.environ.get("PAI_OBS_DAILY_OUTPUT_DIR", "/tmp/paiobs_daily_reports")
DEFAULT_REMOTE_DIR = os.environ.get("PAI_OBS_DAILY_REMOTE_DIR", "tmp")
DEFAULT_RECIPIENT = os.environ.get("PAI_OBS_DAILY_RECIPIENT", "fengchao")
DEFAULT_AI_LIMIT = int(os.environ.get("PAI_OBS_DAILY_AI_LIMIT", "0"))
try:
    DEFAULT_DAILY_AI_WORKERS = max(1, int(os.environ.get("PAI_OBS_DAILY_AI_WORKERS", str(DEFAULT_AI_WORKERS)) or str(DEFAULT_AI_WORKERS)))
except ValueError:
    DEFAULT_DAILY_AI_WORKERS = DEFAULT_AI_WORKERS
DEFAULT_MAX_TASKS = int(os.environ.get("PAI_OBS_DAILY_MAX_TASKS", "1000000"))
DEFAULT_LARK_PROFILE = "enterprise-fengchao"
DEFAULT_QUERY_VIEW_BASE_URL = os.environ.get("PAI_OBS_QUERY_VIEW_BASE_URL", paiobs.DEFAULT_RELEASE_BASE_URL)
DEFAULT_HIGH_DURATION_SECONDS = float(os.environ.get("PAI_OBS_DAILY_HIGH_DURATION_SECONDS", "900"))
DEFAULT_HIGH_TOKEN_THRESHOLD = float(os.environ.get("PAI_OBS_DAILY_HIGH_TOKEN_THRESHOLD", "300000"))
DEFAULT_LOW_SCORE_THRESHOLD = float(os.environ.get("PAI_OBS_DAILY_LOW_SCORE_THRESHOLD", "5"))
FOCUS_SHEET_ORDER = ["失败任务", "用户抱怨任务", "高耗时任务", "高token消耗任务", "低分任务"]
FOCUS_TASK_FIELDS = [
    "关注类型",
    "关注原因",
    "关注阈值",
    "queryid",
    "session_id",
    "task_index",
    "query内容",
    "query链接",
    "answer内容",
    "最终文件产物",
    "请求时间",
    "响应时间",
    "状态",
    "错误信息",
    "成功",
    "入口",
    "调度类型",
    "是否联网",
    "用户ID",
    "用户名",
    "用户机构",
    "机构性质",
    "机构类型",
    "用户类型",
    "用户角色",
    "产品类型",
    "耗时秒",
    "结果评分",
    "查询一级分类",
    "查询二级分类",
    "问题一级分类",
    "问题二级分类",
    "责任人",
    "消耗token",
    "输入token",
    "输出token",
    "缓存token",
    "LLM调用数",
    "消耗credit",
    "低分原因",
    "改进建议",
    "证据",
    "skills",
    "是否有附件",
]
AI_ANALYSIS_FOCUS_FIELDS = [
    "结果评分",
    "查询一级分类",
    "查询二级分类",
    "问题一级分类",
    "问题二级分类",
    "责任人",
    "消耗token",
    "输入token",
    "输出token",
    "缓存token",
    "LLM调用数",
    "消耗credit",
    "低分原因",
    "改进建议",
    "证据",
]
FOCUS_TASK_FIELD_ALIASES = {
    "queryid": ["问题id", "query_id", "question_id", "feedback_question_id"],
}
FOCUS_FAILED_STATUS_TOKENS = {
    "failed",
    "failure",
    "error",
    "timeout",
    "cancelled",
    "canceled",
    "terminated",
    "stopped",
    "失败",
    "异常",
    "超时",
    "取消",
    "已取消",
}
LOW_SCORE_CSV_FIELDS = FOCUS_TASK_FIELDS[3:]


class WorkflowError(RuntimeError):
    pass


def parse_day(value: str) -> datetime:
    try:
        return datetime.strptime(value, "%Y-%m-%d")
    except ValueError as exc:
        raise SystemExit(f"invalid --target-date: {value}, expected YYYY-MM-DD") from exc


def parse_datetime(value: str) -> datetime:
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            parsed = datetime.strptime(value, fmt)
            if fmt == "%Y-%m-%d":
                return parsed
            return parsed
        except ValueError:
            pass
    raise SystemExit(f"invalid datetime: {value}, expected YYYY-MM-DD HH:MM:SS")


def format_dt(value: datetime) -> str:
    return value.strftime("%Y-%m-%d %H:%M:%S")


def workflow_window(args: argparse.Namespace) -> tuple[str, str, str]:
    if args.start_time or args.end_time:
        if not args.start_time or not args.end_time:
            raise SystemExit("--start-time and --end-time must be provided together")
        start = parse_datetime(args.start_time)
        end = parse_datetime(args.end_time)
        if end <= start:
            raise SystemExit("--end-time must be later than --start-time")
        target = args.target_date or start.strftime("%Y-%m-%d")
        return format_dt(start), format_dt(end), target

    tz = ZoneInfo(args.timezone)
    if args.target_date:
        start = parse_day(args.target_date)
    else:
        now = datetime.now(tz)
        start = datetime(now.year, now.month, now.day) - timedelta(days=1)
    end = start + timedelta(days=1)
    return format_dt(start), format_dt(end), start.strftime("%Y-%m-%d")


def compact_window_slug(start_time: str, end_time: str) -> str:
    return (
        f"{start_time}_to_{end_time}"
        .replace("-", "")
        .replace(":", "")
        .replace(" ", "_")
    )


def add_runtime_passthrough(cmd: list[str], args: argparse.Namespace) -> None:
    mapping = [
        ("base_url", "--base-url"),
        ("gateway_profile", "--gateway-profile"),
        ("api_key", "--api-key"),
        ("env", "--env"),
        ("timeout", "--timeout"),
        ("file_auth_token", "--file-auth-token"),
    ]
    for attr, flag in mapping:
        value = getattr(args, attr, None)
        if value not in (None, ""):
            cmd.extend([flag, str(value)])


def add_optional_value(cmd: list[str], args: argparse.Namespace, attr: str, flag: str) -> None:
    value = getattr(args, attr, None)
    if value not in (None, ""):
        cmd.extend([flag, str(value)])


def redact_cmd(cmd: list[str]) -> list[str]:
    redacted = list(cmd)
    secret_flags = {"--api-key", "--file-auth-token", "--ai-api-key"}
    for index, value in enumerate(redacted[:-1]):
        if value in secret_flags:
            redacted[index + 1] = "<redacted>"
    return redacted


def _drain_stream(stream: Any, sink: Any, capture: list[str], echo: bool) -> None:
    try:
        for line in iter(stream.readline, ""):
            capture.append(line)
            if echo:
                sink.write(line)
                sink.flush()
    finally:
        stream.close()


def find_json_payload(text: str) -> Any:
    decoder = json.JSONDecoder()
    for index, char in enumerate(text):
        if char not in "[{":
            continue
        try:
            payload, _end = decoder.raw_decode(text[index:])
            return payload
        except ValueError:
            continue
    raise WorkflowError(f"command did not return JSON: {text.strip()[:500]}")


def run_json_command(cmd: list[str], *, cwd: Path, label: str, verbose: bool = False) -> dict[str, Any]:
    sys.stderr.write(f"[paiobs-daily] {label} start\n")
    sys.stderr.flush()
    started = time.time()
    process = subprocess.Popen(
        cmd,
        cwd=str(cwd),
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        bufsize=1,
        universal_newlines=True,
    )
    stdout_parts: list[str] = []
    stderr_parts: list[str] = []
    stdout_thread = threading.Thread(
        target=_drain_stream,
        args=(process.stdout, sys.stderr, stdout_parts, verbose),
        daemon=True,
    )
    stderr_thread = threading.Thread(
        target=_drain_stream,
        args=(process.stderr, sys.stderr, stderr_parts, True),
        daemon=True,
    )
    stdout_thread.start()
    stderr_thread.start()
    return_code = process.wait()
    stdout_thread.join()
    stderr_thread.join()
    elapsed = time.time() - started
    if return_code != 0:
        stdout_tail = "".join(stdout_parts)[-2000:]
        stderr_tail = "".join(stderr_parts)[-4000:]
        raise WorkflowError(
            f"{label} failed with exit code {return_code}\n"
            f"cmd={redact_cmd(cmd)}\nstdout_tail={stdout_tail}\nstderr_tail={stderr_tail}"
        )
    sys.stderr.write(f"[paiobs-daily] {label} done elapsed={elapsed:.1f}s\n")
    sys.stderr.flush()
    payload = find_json_payload("".join(stdout_parts))
    if not isinstance(payload, dict):
        raise WorkflowError(f"{label} returned non-object JSON")
    return payload


def run_command(cmd: list[str], *, cwd: Path, label: str, verbose: bool = False) -> None:
    sys.stderr.write(f"[paiobs-daily] {label} start\n")
    sys.stderr.flush()
    started = time.time()
    process = subprocess.Popen(
        cmd,
        cwd=str(cwd),
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        bufsize=1,
        universal_newlines=True,
    )
    stdout_parts: list[str] = []
    stderr_parts: list[str] = []
    stdout_thread = threading.Thread(
        target=_drain_stream,
        args=(process.stdout, sys.stderr, stdout_parts, verbose),
        daemon=True,
    )
    stderr_thread = threading.Thread(
        target=_drain_stream,
        args=(process.stderr, sys.stderr, stderr_parts, True),
        daemon=True,
    )
    stdout_thread.start()
    stderr_thread.start()
    return_code = process.wait()
    stdout_thread.join()
    stderr_thread.join()
    elapsed = time.time() - started
    if return_code != 0:
        stdout_tail = "".join(stdout_parts)[-2000:]
        stderr_tail = "".join(stderr_parts)[-4000:]
        raise WorkflowError(
            f"{label} failed with exit code {return_code}\n"
            f"cmd={redact_cmd(cmd)}\nstdout_tail={stdout_tail}\nstderr_tail={stderr_tail}"
        )
    sys.stderr.write(f"[paiobs-daily] {label} done elapsed={elapsed:.1f}s\n")
    sys.stderr.flush()


def build_stats_command(
    args: argparse.Namespace,
    *,
    start_time: str,
    end_time: str,
    quick_stats_md: Path,
    quick_stats_json: Path,
) -> list[str]:
    cmd = [sys.executable, str(SCRIPT_DIR / "paiobs_quick_stats_report.py")]
    add_runtime_passthrough(cmd, args)
    cmd.extend(
        [
            "--start-time",
            start_time,
            "--end-time",
            end_time,
            "--report-md",
            str(quick_stats_md),
            "--json-output",
            str(quick_stats_json),
            "--breakdowns",
            args.breakdowns,
            "--detail-dimensions",
            args.task_dimensions,
            "--max-tasks",
            str(args.max_tasks),
            "--limit",
            str(args.stats_limit),
            "--display-top-n",
            str(args.top_n),
            "--detail-top-n",
            str(args.top_n),
            "--format",
            "json",
        ]
    )
    return cmd


def build_ai_command(
    args: argparse.Namespace,
    *,
    start_time: str,
    end_time: str,
    analysis_jsonl: Path,
) -> list[str]:
    limit = args.ai_limit if args.ai_limit and args.ai_limit > 0 else args.max_tasks
    cmd = [sys.executable, str(SCRIPT_DIR / "paiobs_ai_analysis.py")]
    add_runtime_passthrough(cmd, args)
    cmd.extend(
        [
            "analyze-batch",
            "--start-time",
            start_time,
            "--end-time",
            end_time,
            "--limit",
            str(limit),
            "--complete-search",
            "--context-profile",
            args.context_profile,
            "--output",
            str(analysis_jsonl),
            "--workers",
            str(args.ai_workers),
            "--progress",
            "--ended-only",
        ]
    )
    if args.include_process:
        cmd.append("--include-process")
    if args.include_raw_context:
        cmd.append("--include-raw-context")
    if args.include_file_previews:
        cmd.append("--include-file-previews")
    if not args.shortcut_failed:
        cmd.append("--no-shortcut-failed")
    for attr, flag in [
        ("ai_provider", "--ai-provider"),
        ("ai_base_url", "--ai-base-url"),
        ("ai_api_key", "--ai-api-key"),
        ("ai_model", "--ai-model"),
        ("ai_fallback_models", "--ai-fallback-models"),
        ("ai_timeout", "--ai-timeout"),
        ("ai_retries", "--ai-retries"),
        ("temperature", "--temperature"),
        ("max_tokens", "--max-tokens"),
    ]:
        add_optional_value(cmd, args, attr, flag)
    return cmd


def fetch_focus_tasks_report(args: argparse.Namespace, *, start_time: str, end_time: str) -> dict[str, Any]:
    client = paiobs.build_client(args)
    body = {
        "env": client.env,
        "filters": {
            "start_time": start_time,
            "end_time": end_time,
        },
        "max_items": args.max_tasks,
        "high_duration_seconds": args.high_duration_seconds,
        "high_token_threshold": args.high_token_threshold,
        "query_view_base_url": DEFAULT_QUERY_VIEW_BASE_URL,
        "search_workers": args.focus_search_workers,
        "detail_workers": args.focus_detail_workers,
        "search_page_limit": args.focus_search_page_limit,
        "search_slice_minutes": args.focus_search_slice_minutes,
        "min_slice_seconds": args.focus_min_slice_seconds,
        "sheet_limit": args.focus_sheet_limit,
    }
    sys.stderr.write("[paiobs-daily] focus tasks gateway report start\n")
    sys.stderr.flush()
    started = time.time()
    payload = client.request("POST", "/reports/focus-tasks", body=body)
    sys.stderr.write(f"[paiobs-daily] focus tasks gateway report done elapsed={time.time() - started:.1f}s\n")
    sys.stderr.flush()
    if not isinstance(payload, dict) or payload.get("schema_version") != "focus-tasks-report/v1":
        raise WorkflowError("focus tasks gateway report returned unexpected payload")
    return payload


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    if not path.is_file():
        return items
    with path.open("r", encoding="utf-8") as fh:
        for line in fh:
            text = line.strip()
            if not text:
                continue
            try:
                value = json.loads(text)
            except json.JSONDecodeError:
                continue
            if isinstance(value, dict):
                items.append(value)
    return items


def score_value(item: dict[str, Any]) -> float:
    for key in ("result_score", "overall_score"):
        value = item.get(key)
        if value in (None, ""):
            continue
        try:
            return float(value)
        except (TypeError, ValueError):
            continue
    return 0.0


def low_score_items(items: list[dict[str, Any]], threshold: float = DEFAULT_LOW_SCORE_THRESHOLD) -> list[dict[str, Any]]:
    return [item for item in items if score_value(item) < threshold]


def item_metadata(item: dict[str, Any]) -> dict[str, Any]:
    return item.get("task_metadata") if isinstance(item.get("task_metadata"), dict) else {}


def first_value(item: dict[str, Any], *keys: str) -> Any:
    metadata = item_metadata(item)
    for key in keys:
        value = item.get(key)
        if value not in (None, "", [], {}):
            return value
        value = metadata.get(key)
        if value not in (None, "", [], {}):
            return value
    return ""


def join_cell(value: Any) -> str:
    if isinstance(value, list):
        return "；".join(str(item) for item in value if item not in (None, "", [], {}))
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return "" if value is None else str(value)


def csv_cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (list, dict)):
        return join_cell(value)
    return value


def has_value(value: Any) -> bool:
    return value not in (None, "", [], {})


def user_role_label(value: Any) -> str:
    return str(paiobs.normalize_dimension_value("user_role", value)) if value not in (None, "") else ""


def final_output_files(item: dict[str, Any]) -> list[dict[str, Any]]:
    value = first_value(item, "final_output_files")
    if isinstance(value, list):
        return [entry for entry in value if isinstance(entry, dict)]
    return []


def first_file_content_url(files: list[dict[str, Any]]) -> str:
    for file_item in files:
        link = str(file_item.get("content_url") or "").strip()
        if link:
            return link
    for file_item in files:
        link = str(file_item.get("preview_url") or "").strip()
        if link:
            return link
    return ""


def boolish(value: Any) -> bool | None:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if value == 1:
            return True
        if value == 0:
            return False
    text = str(value).strip().lower()
    if text in {"true", "1", "yes", "y", "scheduled"}:
        return True
    if text in {"false", "0", "no", "n", "manual"}:
        return False
    return None


def scheduled_label(value: Any) -> str:
    text = str(value).strip()
    if text in {"scheduled", "manual"}:
        return text
    flag = boolish(value)
    if flag is True:
        return "scheduled"
    if flag is False:
        return "manual"
    return csv_cell(value)


def web_search_label(value: Any) -> str:
    text = str(value).strip().lower()
    if text in {"true", "false"}:
        return text
    flag = boolish(value)
    if flag is True:
        return "true"
    if flag is False:
        return "false"
    return csv_cell(value)


def focus_row_is_failed(row: dict[str, Any]) -> bool:
    if boolish(row.get("成功")) is False:
        return True
    status = str(row.get("状态") or "").strip().lower()
    return any(token in status for token in FOCUS_FAILED_STATUS_TOKENS)


def query_view_url(item: dict[str, Any]) -> str:
    query_id = str(first_value(item, "query_id", "queryid", "question_id", "feedback_question_id") or "").strip()
    if not query_id:
        return ""
    base_url = (DEFAULT_QUERY_VIEW_BASE_URL or paiobs.DEFAULT_RELEASE_BASE_URL).rstrip("/")
    return f"{base_url}/?{urlencode({'queryid': query_id})}"


def focus_row_from_analysis_item(
    item: dict[str, Any],
    *,
    focus_type: str,
    focus_reason: str,
    focus_threshold: str,
) -> dict[str, Any]:
    output_files = final_output_files(item)
    issue_level2 = normalize_issue_level2(item.get("issue_level2"))
    issue_level1 = issue_level1_for_level2(issue_level2, item.get("issue_level1"))
    owner = issue_owner_for_level2(issue_level2)
    row = {
        "关注类型": focus_type,
        "关注原因": focus_reason,
        "关注阈值": focus_threshold,
        "queryid": first_value(item, "query_id", "queryid", "question_id", "feedback_question_id"),
        "session_id": first_value(item, "session_id"),
        "task_index": first_value(item, "task_index"),
        "query内容": first_value(item, "question"),
        "query链接": query_view_url(item),
        "answer内容": first_value(item, "answer", "final_answer_excerpt"),
        "最终文件产物": first_file_content_url(output_files),
        "请求时间": first_value(item, "request_time"),
        "响应时间": first_value(item, "response_time"),
        "状态": first_value(item, "status"),
        "错误信息": first_value(item, "error"),
        "成功": first_value(item, "success"),
        "入口": first_value(item, "entry_scene"),
        "调度类型": scheduled_label(first_value(item, "scheduled")),
        "是否联网": web_search_label(first_value(item, "is_web_search")),
        "用户ID": first_value(item, "user_id"),
        "用户名": first_value(item, "user_name"),
        "用户机构": first_value(item, "user_institution", "institution"),
        "机构性质": first_value(item, "institution_nature"),
        "机构类型": first_value(item, "inst_type"),
        "用户类型": first_value(item, "user_type"),
        "用户角色": user_role_label(first_value(item, "user_role")),
        "产品类型": first_value(item, "product_type"),
        "耗时秒": first_value(item, "duration_seconds", "duration", "dur", "total_elapsed", "average_time"),
        "结果评分": csv_cell(item.get("result_score") if item.get("result_score") not in (None, "") else item.get("overall_score")),
        "查询一级分类": chinese_label(item.get("query_major_category"), QUERY_MAJOR_CATEGORY_LABELS, "其他"),
        "查询二级分类": chinese_label(item.get("query_minor_category"), QUERY_MINOR_CATEGORY_LABELS, "不明确或非金融问题"),
        "问题一级分类": issue_level1,
        "问题二级分类": issue_level2,
        "责任人": owner,
        "消耗token": csv_cell(item.get("token_total")),
        "输入token": csv_cell(item.get("prompt_tokens")),
        "输出token": csv_cell(item.get("completion_tokens")),
        "缓存token": csv_cell(item.get("cached_tokens")),
        "LLM调用数": csv_cell(item.get("llm_call_count")),
        "消耗credit": csv_cell(item.get("credit_total")),
        "低分原因": item.get("low_score_reason") or "",
        "改进建议": item.get("improvement_suggestion") or "",
        "证据": join_cell(item.get("evidence")),
        "skills": join_cell(first_value(item, "skills")),
        "是否有附件": first_value(item, "has_attachments"),
    }
    return {field: row.get(field, "") for field in FOCUS_TASK_FIELDS}


def low_score_rows(items: list[dict[str, Any]], *, threshold: float) -> list[dict[str, Any]]:
    rows = []
    for item in low_score_items(items, threshold=threshold):
        score = score_value(item)
        row = focus_row_from_analysis_item(
            item,
            focus_type="低分任务",
            focus_reason=item.get("low_score_reason") or f"结果评分 {score:g} 低于关注阈值",
            focus_threshold=f"结果评分 < {threshold:g}",
        )
        if focus_row_is_failed(row):
            continue
        rows.append(row)
    rows.sort(key=lambda row: (paiobs.safe_number(row.get("结果评分")) if hasattr(paiobs, "safe_number") else None) or 0)
    return rows


def write_low_score_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=LOW_SCORE_CSV_FIELDS)
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field, "") for field in LOW_SCORE_CSV_FIELDS})


def focus_sheets_from_gateway(focus_payload: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    result: dict[str, list[dict[str, Any]]] = {sheet: [] for sheet in FOCUS_SHEET_ORDER}
    sheets = focus_payload.get("sheets") if isinstance(focus_payload.get("sheets"), dict) else {}
    for sheet_name in FOCUS_SHEET_ORDER:
        sheet = sheets.get(sheet_name) if isinstance(sheets.get(sheet_name), dict) else {}
        items = sheet.get("items") if isinstance(sheet.get("items"), list) else []
        result[sheet_name] = [
            normalize_focus_row_issue_fields({field: focus_field_value(item, field) for field in FOCUS_TASK_FIELDS})
            for item in items
            if isinstance(item, dict)
        ]
    return result


def focus_field_value(item: dict[str, Any], field: str) -> Any:
    value = item.get(field, "")
    if value not in (None, "", [], {}):
        pass  # found
    else:
        for alias in FOCUS_TASK_FIELD_ALIASES.get(field, []):
            value = item.get(alias, "")
            if value not in (None, "", [], {}):
                break
        else:
            value = "" if value is None else value
    # Normalize user_role from numeric to Chinese label
    if field == "用户角色":
        value = user_role_label(value)
    return value


def normalize_focus_row_issue_fields(row: dict[str, Any]) -> dict[str, Any]:
    issue_level2 = normalize_issue_level2(row.get("问题二级分类") or row.get("issue_level2"))
    row["问题二级分类"] = issue_level2
    row["问题一级分类"] = issue_level1_for_level2(issue_level2, row.get("问题一级分类") or row.get("issue_level1"))
    row["责任人"] = issue_owner_for_level2(issue_level2)
    return {field: row.get(field, "") for field in FOCUS_TASK_FIELDS}


def normalized_identity_value(value: Any) -> str:
    return str(value or "").strip()


def normalized_task_index(value: Any) -> str:
    text = normalized_identity_value(value)
    if not text:
        return ""
    try:
        number = float(text)
    except ValueError:
        return text
    if number.is_integer():
        return str(int(number))
    return text


def focus_row_identity_keys(row: dict[str, Any]) -> list[tuple[str, str]]:
    query_id = normalized_identity_value(row.get("queryid"))
    session_id = normalized_identity_value(row.get("session_id"))
    task_index = normalized_task_index(row.get("task_index"))
    keys: list[tuple[str, str]] = []
    if session_id and task_index and query_id:
        keys.append(("session_task_query", f"{session_id}:{task_index}:{query_id}"))
    if query_id:
        keys.append(("queryid", query_id))
    if session_id and task_index:
        keys.append(("session_task", f"{session_id}:{task_index}"))
    return keys


def analysis_item_identity_keys(item: dict[str, Any]) -> list[tuple[str, str]]:
    query_id = normalized_identity_value(
        first_value(item, "query_id", "queryid", "question_id", "feedback_question_id")
    )
    session_id = normalized_identity_value(first_value(item, "session_id"))
    task_index = normalized_task_index(first_value(item, "task_index"))
    keys: list[tuple[str, str]] = []
    if session_id and task_index and query_id:
        keys.append(("session_task_query", f"{session_id}:{task_index}:{query_id}"))
    if query_id:
        keys.append(("queryid", query_id))
    if session_id and task_index:
        keys.append(("session_task", f"{session_id}:{task_index}"))
    return keys


def analysis_focus_row(item: dict[str, Any]) -> dict[str, Any]:
    return focus_row_from_analysis_item(
        item,
        focus_type="",
        focus_reason="",
        focus_threshold="",
    )


def build_analysis_focus_index(items: list[dict[str, Any]]) -> dict[tuple[str, str], dict[str, Any]]:
    index: dict[tuple[str, str], dict[str, Any]] = {}
    for item in items:
        row = analysis_focus_row(item)
        for key in analysis_item_identity_keys(item):
            index.setdefault(key, row)
    return index


def enrich_focus_row_with_ai(row: dict[str, Any], ai_row: dict[str, Any]) -> dict[str, Any]:
    enriched = dict(row)
    for field in AI_ANALYSIS_FOCUS_FIELDS:
        value = ai_row.get(field)
        if has_value(value):
            enriched[field] = value
    return normalize_focus_row_issue_fields(enriched)


def enrich_focus_sheets_with_ai(
    sheets: dict[str, list[dict[str, Any]]],
    analysis_items: list[dict[str, Any]],
) -> tuple[dict[str, list[dict[str, Any]]], dict[str, Any]]:
    analysis_index = build_analysis_focus_index(analysis_items)
    enriched_sheets: dict[str, list[dict[str, Any]]] = {sheet: [] for sheet in FOCUS_SHEET_ORDER}
    matched_by_sheet: dict[str, int] = {}
    unmatched_by_sheet: dict[str, int] = {}
    total_rows = 0
    matched_rows = 0
    for sheet_name in FOCUS_SHEET_ORDER:
        matched = 0
        unmatched = 0
        rows = []
        for row in sheets.get(sheet_name, []):
            total_rows += 1
            ai_row = None
            for key in focus_row_identity_keys(row):
                ai_row = analysis_index.get(key)
                if ai_row is not None:
                    break
            if ai_row is None:
                unmatched += 1
                rows.append(normalize_focus_row_issue_fields({field: row.get(field, "") for field in FOCUS_TASK_FIELDS}))
                continue
            matched += 1
            matched_rows += 1
            rows.append(enrich_focus_row_with_ai(row, ai_row))
        enriched_sheets[sheet_name] = rows
        matched_by_sheet[sheet_name] = matched
        unmatched_by_sheet[sheet_name] = unmatched
    return enriched_sheets, {
        "total_rows": total_rows,
        "matched_rows": matched_rows,
        "unmatched_rows": total_rows - matched_rows,
        "matched_by_sheet": matched_by_sheet,
        "unmatched_by_sheet": unmatched_by_sheet,
    }


def _compute_nonempty_fields(sheets: dict[str, list[dict[str, Any]]]) -> set[str]:
    """Return the set of fields that have at least one non-empty value across all sheets."""
    nonempty: set[str] = set()
    for rows in sheets.values():
        for row in rows:
            for field in FOCUS_TASK_FIELDS:
                if field in nonempty:
                    continue
                value = row.get(field, "")
                if value not in (None, "", [], {}):
                    nonempty.add(field)
    return nonempty


# Fields that must always be included even if all values are empty.
ALWAYS_KEEP_FIELDS = {
    "关注类型", "关注原因", "queryid", "session_id", "task_index",
    "query内容", "query链接", "answer内容", "最终文件产物",
    "请求时间", "响应时间", "状态", "成功",
    "入口", "调度类型", "是否联网",
    "用户名", "用户机构", "用户类型", "用户角色", "产品类型",
    "结果评分", "查询一级分类", "查询二级分类",
    "问题一级分类", "问题二级分类", "责任人",
    "低分原因", "改进建议",
}


def write_focus_task_workbook(path: Path, sheets: dict[str, list[dict[str, Any]]]) -> None:
    if Workbook is None:
        raise WorkflowError("openpyxl is required to write the focus task workbook")
    path.parent.mkdir(parents=True, exist_ok=True)
    nonempty = _compute_nonempty_fields(sheets)
    active_fields = [f for f in FOCUS_TASK_FIELDS if f in nonempty or f in ALWAYS_KEEP_FIELDS]
    workbook = Workbook()
    for index, sheet_name in enumerate(FOCUS_SHEET_ORDER):
        worksheet = workbook.active if index == 0 else workbook.create_sheet()
        worksheet.title = sheet_name[:31]
        worksheet.append(active_fields)
        for row in sheets.get(sheet_name, []):
            worksheet.append([csv_cell(row.get(field, "")) for field in active_fields])
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for column_cells in worksheet.columns:
            header = str(column_cells[0].value or "")
            max_len = max([len(str(cell.value or "")) for cell in column_cells[: min(len(column_cells), 50)]], default=len(header))
            width = min(max(max_len + 2, len(header) + 2, 10), 60)
            worksheet.column_dimensions[column_cells[0].column_letter].width = width
    workbook.save(path)


def markdown_table(rows: list[list[Any]], headers: list[str]) -> str:
    lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join("---" for _ in headers) + " |",
    ]
    for row in rows:
        lines.append("| " + " | ".join(("" if value is None else str(value)).replace("\n", " ") for value in row) + " |")
    return "\n".join(lines)


def render_daily_report(
    *,
    start_time: str,
    end_time: str,
    quick_stats_md: Path,
    analysis_items: list[dict[str, Any]],
    focus_task_file: Path,
    focus_sheets: dict[str, list[dict[str, Any]]],
    low_score_threshold: float,
) -> str:
    low_sheet_rows = focus_sheets.get("低分任务", [])
    avg_score = round(sum(score_value(item) for item in analysis_items) / len(analysis_items), 2) if analysis_items else ""
    low_rows = [
        [
            row.get("请求时间") or "",
            row.get("用户名") or "",
            row.get("用户机构") or "",
            row.get("query内容") or "",
            row.get("queryid") or "",
            row.get("session_id") or "",
            row.get("task_index") or "",
            row.get("结果评分") or "",
            row.get("问题一级分类") or "",
            row.get("问题二级分类") or "",
            row.get("低分原因") or "",
        ]
        for row in low_sheet_rows[:20]
    ]
    quick_stats_text = quick_stats_md.read_text(encoding="utf-8") if quick_stats_md.is_file() else "未生成聚合统计报告。"
    return "\n".join(
        [
            "# PaiWork 完整日报",
            "",
            f"- 统计窗口：{start_time} ~ {end_time} Asia/Shanghai",
            "- 报告结构：聚合统计 + 已结束任务 AI 分析 + 关注任务表格",
            "",
            "## 1. 聚合统计",
            "",
            quick_stats_text,
            "",
            "## 2. AI 分析",
            "",
            markdown_table(
                [
                    ["已结束任务 AI 标注数", len(analysis_items)],
                    ["低分任务数", len(low_sheet_rows)],
                    ["平均结果评分", avg_score],
                    ["关注任务表格", str(focus_task_file)],
                ],
                ["指标", "数值"],
            ),
            "",
            "## 3. 关注任务表格",
            "",
            markdown_table(
                [
                    ["失败任务", len(focus_sheets.get("失败任务", [])), "已结束任务且 success=false 或状态为 failed/error/timeout/cancelled"],
                    ["用户抱怨任务", len(focus_sheets.get("用户抱怨任务", [])), "已结束任务且 query 内容命中用户抱怨关键词"],
                    ["高耗时任务", len(focus_sheets.get("高耗时任务", [])), "已结束且非失败任务，达到高耗时阈值"],
                    ["高token消耗任务", len(focus_sheets.get("高token消耗任务", [])), "已结束且非失败任务，达到高 token 阈值"],
                    ["低分任务", len(focus_sheets.get("低分任务", [])), f"已结束且非失败任务，结果评分 < {low_score_threshold:g}"],
                ],
                ["sheet", "任务数", "口径"],
            ),
            "",
            "## 4. 低分样本",
            "",
            markdown_table(low_rows, ["请求时间", "用户", "机构", "query内容", "queryid", "session_id", "task", "结果评分", "问题一级", "问题二级", "原因"])
            if low_rows
            else "本期无低分样本。",
            "",
        ]
    ).rstrip() + "\n"


def build_publish_command(
    args: argparse.Namespace,
    *,
    start_time: str,
    end_time: str,
    target_date: str,
    report_md: Path,
    focus_task_file: Path,
    report_payload: dict[str, Any],
) -> list[str]:
    doc_path = args.doc_path or f"{args.remote_dir}/{target_date}_paiwork_daily_report.md"
    bitable_path = args.bitable_path or f"{args.remote_dir}/{target_date}_paiwork_focus_tasks.bitable.xlsx"
    message = args.message
    if not message:
        analysis_count = report_payload.get("analysis_count", "")
        low_count = report_payload.get("low_score_count", "")
        focus_counts = report_payload.get("focus_task_counts") if isinstance(report_payload.get("focus_task_counts"), dict) else {}
        message = "\n".join(
            [
                "PaiWork 完整日报已生成。",
                f"时间范围：{start_time} 至 {end_time}",
                f"已结束任务 AI 标注数：{analysis_count}；低分任务数：{low_count}",
                f"关注任务：失败 {focus_counts.get('失败任务', 0)}；用户抱怨 {focus_counts.get('用户抱怨任务', 0)}；高耗时 {focus_counts.get('高耗时任务', 0)}；高token {focus_counts.get('高token消耗任务', 0)}；低分 {focus_counts.get('低分任务', 0)}",
                "报告：{doc_url}",
                "关注任务表：{bitable_url}",
            ]
        )
    cmd = [
        sys.executable,
        str(SCRIPT_DIR / "paiobs_lark.py"),
        "publish",
        "--report-md",
        str(report_md),
        "--focus-task-file",
        str(focus_task_file),
        "--doc-path",
        doc_path,
        "--bitable-path",
        bitable_path,
        "--recipient",
        args.recipient,
        "--message",
        message,
        "--lark-profile",
        args.lark_profile,
        "--send-as",
        args.send_as,
        "--format",
        "json",
    ]
    if args.send:
        cmd.append("--send")
    if args.recipient_user_id:
        cmd.extend(["--recipient-user-id", args.recipient_user_id])
    if args.dry_run:
        cmd.append("--dry-run")
    return cmd


def write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def cmd_run(args: argparse.Namespace) -> None:
    start_time, end_time, target_date = workflow_window(args)
    output_dir = Path(args.output_dir).expanduser().resolve() / compact_window_slug(start_time, end_time)
    output_dir.mkdir(parents=True, exist_ok=True)
    stem = f"paiwork_daily_{compact_window_slug(start_time, end_time)}"
    report_md = output_dir / f"{stem}.md"
    quick_stats_md = output_dir / f"{stem}_aggregate_stats.md"
    quick_stats_json = output_dir / f"{stem}_aggregate_stats.json"
    focus_task_file = output_dir / f"{stem}_focus_tasks.xlsx"
    low_score_csv = output_dir / f"{stem}_low_scores.csv"
    analysis_jsonl = output_dir / f"{stem}_analysis.jsonl"
    manifest_path = Path(args.manifest).expanduser().resolve() if args.manifest else output_dir / f"{stem}_manifest.json"

    stats_cmd = build_stats_command(
        args,
        start_time=start_time,
        end_time=end_time,
        quick_stats_md=quick_stats_md,
        quick_stats_json=quick_stats_json,
    )
    stats_payload = run_json_command(stats_cmd, cwd=SKILL_DIR, label="aggregate stats", verbose=args.verbose)
    focus_payload = fetch_focus_tasks_report(args, start_time=start_time, end_time=end_time)

    ai_cmd = build_ai_command(args, start_time=start_time, end_time=end_time, analysis_jsonl=analysis_jsonl)
    run_command(ai_cmd, cwd=SKILL_DIR, label="ai analysis", verbose=args.verbose)
    analysis_items = read_jsonl(analysis_jsonl)
    focus_sheets = focus_sheets_from_gateway(focus_payload)
    low_rows = low_score_rows(analysis_items, threshold=args.low_score_threshold)
    focus_sheets["低分任务"] = low_rows
    focus_sheets, enrichment_summary = enrich_focus_sheets_with_ai(focus_sheets, analysis_items)
    write_focus_task_workbook(focus_task_file, focus_sheets)
    write_low_score_csv(low_score_csv, low_rows)
    report_md.write_text(
        render_daily_report(
            start_time=start_time,
            end_time=end_time,
            quick_stats_md=quick_stats_md,
            analysis_items=analysis_items,
            focus_task_file=focus_task_file,
            focus_sheets=focus_sheets,
            low_score_threshold=args.low_score_threshold,
        ),
        encoding="utf-8",
    )
    focus_counts = {sheet: len(focus_sheets.get(sheet, [])) for sheet in FOCUS_SHEET_ORDER}
    report_payload = {
        "schema_version": "daily-report-output/v1",
        "report_md": str(report_md),
        "quick_stats_md": str(quick_stats_md),
        "quick_stats_json": str(quick_stats_json),
        "focus_task_file": str(focus_task_file),
        "analysis_jsonl": str(analysis_jsonl),
        "low_score_csv": str(low_score_csv),
        "analysis_count": len(analysis_items),
        "low_score_count": len(low_rows),
        "focus_task_counts": focus_counts,
        "focus_task_ai_enrichment": enrichment_summary,
        "focus_tasks": focus_payload,
        "thresholds": {
            "high_duration_seconds": args.high_duration_seconds,
            "high_token_threshold": args.high_token_threshold,
            "low_score_threshold": args.low_score_threshold,
        },
        "aggregate": stats_payload,
    }

    publish_payload: dict[str, Any] = {}
    if args.publish:
        publish_cmd = build_publish_command(
            args,
            start_time=start_time,
            end_time=end_time,
            target_date=target_date,
            report_md=report_md,
            focus_task_file=focus_task_file,
            report_payload=report_payload,
        )
        publish_payload = run_json_command(publish_cmd, cwd=SKILL_DIR, label="feishu publish", verbose=args.verbose)

    manifest = {
        "schema_version": "paiobs-daily-workflow/v1",
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "target_date": target_date,
        "timezone": args.timezone,
        "window": {
            "start_time": start_time,
            "end_time": end_time,
        },
        "ai_limit": args.ai_limit,
        "max_tasks": args.max_tasks,
        "recipient": args.recipient,
        "sent": bool(args.publish and args.send and not args.dry_run),
        "dry_run": bool(args.dry_run),
        "outputs": {
            "output_dir": str(output_dir),
            "report_md": str(report_md),
            "quick_stats_md": str(quick_stats_md),
            "quick_stats_json": str(quick_stats_json),
            "focus_task_file": str(focus_task_file),
            "low_score_csv": str(low_score_csv),
            "analysis_jsonl": str(analysis_jsonl),
            "manifest": str(manifest_path),
        },
        "report": report_payload,
        "publish": publish_payload,
    }
    write_json(manifest_path, manifest)
    sys.stdout.write(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Run the PaiWork daily report workflow")
    parser.add_argument("--target-date", default="", help="Local day to analyze, YYYY-MM-DD. Defaults to yesterday.")
    parser.add_argument("--start-time", default="", help="Override window start, YYYY-MM-DD HH:MM:SS")
    parser.add_argument("--end-time", default="", help="Override window end, YYYY-MM-DD HH:MM:SS")
    parser.add_argument("--timezone", default=DEFAULT_TIMEZONE)
    parser.add_argument("--output-dir", default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--manifest", default="")
    parser.add_argument("--remote-dir", default=DEFAULT_REMOTE_DIR)
    parser.add_argument("--doc-path", default="")
    parser.add_argument("--bitable-path", default="")
    parser.add_argument("--publish", action=argparse.BooleanOptionalAction, default=True)
    parser.add_argument("--send", action=argparse.BooleanOptionalAction, default=True)
    parser.add_argument("--recipient", default=DEFAULT_RECIPIENT)
    parser.add_argument("--recipient-user-id", default="")
    parser.add_argument("--message", default="")
    parser.add_argument("--lark-profile", default=os.environ.get("PAI_OBS_LARK_PROFILE", DEFAULT_LARK_PROFILE))
    parser.add_argument("--send-as", choices=["auto", "user", "bot"], default=os.environ.get("PAI_OBS_LARK_SEND_AS", "auto"))
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--verbose", action="store_true", help="Echo child command stdout as well as stderr")

    parser.add_argument("--ai-limit", type=int, default=DEFAULT_AI_LIMIT, help="0 means AI-label all matching ended tasks")
    parser.add_argument("--ai-workers", type=int, default=DEFAULT_DAILY_AI_WORKERS, help="Concurrent AI judging workers passed to paiobs_ai_analysis.py")
    parser.add_argument("--max-tasks", type=int, default=DEFAULT_MAX_TASKS)
    parser.add_argument("--stats-limit", type=int, default=500)
    parser.add_argument("--top-n", type=int, default=10)
    parser.add_argument("--high-duration-seconds", type=float, default=DEFAULT_HIGH_DURATION_SECONDS)
    parser.add_argument("--high-token-threshold", type=float, default=DEFAULT_HIGH_TOKEN_THRESHOLD)
    parser.add_argument("--low-score-threshold", type=float, default=DEFAULT_LOW_SCORE_THRESHOLD)
    parser.add_argument("--focus-search-workers", type=int, default=8)
    parser.add_argument("--focus-detail-workers", type=int, default=8)
    parser.add_argument("--focus-search-page-limit", type=int, default=100)
    parser.add_argument("--focus-search-slice-minutes", type=float, default=3.0)
    parser.add_argument("--focus-min-slice-seconds", type=float, default=1.0)
    parser.add_argument("--focus-sheet-limit", type=int, default=0, help="0 keeps all matching focus rows returned by the gateway")
    parser.add_argument(
        "--breakdowns",
        default="status,entry_scene,scheduled,is_web_search,web_query_language,end_type,user_role,user_type,institution,product_type",
    )
    parser.add_argument(
        "--task-dimensions",
        default="skill,tool_name,model,data_source_type,file_type",
    )
    parser.add_argument("--fetch-duration-stats", action=argparse.BooleanOptionalAction, default=True, help="Deprecated; aggregate stats follows quickstat template.")
    parser.add_argument("--fetch-token-stats", action=argparse.BooleanOptionalAction, default=True, help="Deprecated; token totals are included by aggregate detail task-stats.")
    parser.add_argument("--fetch-credit-stats", action=argparse.BooleanOptionalAction, default=True, help="Deprecated; credit remains in focus task rows, not aggregate stats.")

    parser.add_argument("--context-profile", default="summary", choices=["summary", "context", "qa", "full"])
    parser.add_argument("--include-process", action="store_true")
    parser.add_argument("--include-raw-context", action="store_true")
    parser.add_argument("--include-file-previews", action="store_true")
    parser.add_argument("--shortcut-failed", action=argparse.BooleanOptionalAction, default=True)
    parser.add_argument("--ai-provider", choices=["gateway", "openai-compatible", "none"], default="")
    parser.add_argument("--ai-base-url", default="")
    parser.add_argument("--ai-api-key", default="")
    parser.add_argument("--ai-model", default="")
    parser.add_argument("--ai-fallback-models", default="")
    parser.add_argument("--ai-timeout", type=float, default=None)
    parser.add_argument("--ai-retries", type=int, default=None)
    parser.add_argument("--temperature", type=float, default=None)
    parser.add_argument("--max-tokens", type=int, default=None)

    parser.add_argument("--base-url", default=None)
    parser.add_argument("--gateway-profile", default=None, help="Gateway profile: release/prod -> 30100, local/test/debug -> localhost:6193")
    parser.add_argument("--api-key", default=None)
    parser.add_argument("--env", default=None)
    parser.add_argument("--timeout", default=None)
    parser.add_argument("--file-auth-token", default=None)
    parser.set_defaults(func=cmd_run)
    return parser


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    try:
        args.func(args)
        return 0
    except WorkflowError as exc:
        sys.stderr.write(f"ERROR: {exc}\n")
        return 1
    except KeyboardInterrupt:
        sys.stderr.write("Interrupted\n")
        return 130


if __name__ == "__main__":
    raise SystemExit(main())
